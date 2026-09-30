import tempfile
import unittest
from pathlib import Path

from logic import db_automation, multi_schema_combiner, table_comparator, workfile_generator


class RegressionTests(unittest.TestCase):
    def test_trigger_debug_calls_are_commented(self):
        raw = "CREATE TRIGGER t BEFORE INSERT ON x FOR EACH ROW BEGIN\n"
        raw += "\tCALL sys_log_debug_sp('x');\n\tCALL sys_log_debug_v9_sp('x');\n"
        raw += "\t-- CALL sys_log_debug_sp('already');\n\tCALL sys_log_error_sp('error');\nEND"
        out = db_automation.process_single_object(raw, 2, 'V20260916', '16 Sep 2026')[0]
        self.assertIn("\t-- CALL sys_log_debug_sp('x');", out)
        self.assertIn("\t-- CALL sys_log_debug_v9_sp('x');", out)
        self.assertIn("\tCALL sys_log_error_sp('error');", out)
        self.assertNotIn('-- -- CALL', out)

    def _convert_function(self, call_line: str) -> str:
        sql = (
            "CREATE FUNCTION `sample_fn` ()\n"
            "RETURNS INT\n"
            "BEGIN\n"
            f"{call_line}\n"
            "    RETURN 1;\n"
            "END;\n"
        )
        result = db_automation.process_single_object(
            sql,
            sequence_num=1,
            date_str="V20260423",
            timestamp_str="23 Apr 2026",
            developer_name="Tester",
        )
        self.assertIsNotNone(result)
        return result[0]

    def test_db_automation_comments_debug_v0_to_v9_and_base(self):
        for proc_name in ["sys_log_debug_sp", "sys_log_debug_v0_sp", "sys_log_debug_v9_sp"]:
            converted = self._convert_function(f"    CALL {proc_name}('x');")
            self.assertIn(f"-- CALL {proc_name}('x');", converted)

    def test_db_automation_does_not_comment_debug_v10(self):
        converted = self._convert_function("    CALL sys_log_debug_v10_sp('x');")
        self.assertIn("CALL sys_log_debug_v10_sp('x');", converted)
        self.assertNotIn("-- CALL sys_log_debug_v10_sp('x');", converted)

    def test_db_automation_blank_description_preserves_blank_header(self):
        converted = self._convert_function("    RETURN 1;")
        self.assertIn("-- Description            :     \n", converted)

    def test_db_automation_default_developer_is_blank(self):
        sql = (
            "CREATE PROCEDURE `sample_sp` ()\n"
            "BEGIN\n"
            "    SELECT 1;\n"
            "END;\n"
        )
        result = db_automation.process_single_object(
            sql,
            sequence_num=1,
            date_str="V20260423",
            timestamp_str="23 Apr 2026",
        )
        self.assertIsNotNone(result)
        self.assertIn("-- Created/Modified By    :    \n", result[0])

    def test_db_automation_description_flows_to_pasted_and_folder_outputs(self):
        sql = (
            "CREATE PROCEDURE `sample_sp` ()\n"
            "BEGIN\n"
            "    SELECT 1;\n"
            "END;\n"
        )

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)

            paste_output_dir = tmp_path / "paste-output"
            pasted = db_automation.process_pasted_to_files(
                sql,
                str(paste_output_dir),
                starting_seq=1,
                developer_name="Tester",
                description="Create sample automation",
            )
            self.assertEqual(pasted["errors"], [])
            pasted_file = paste_output_dir / pasted["processed"][0]["output"]
            pasted_content = pasted_file.read_text(encoding="utf-8")
            self.assertIn(
                "-- Description            :    Create sample automation",
                pasted_content,
            )

            folder_input_dir = tmp_path / "folder-input"
            folder_output_dir = tmp_path / "folder-output"
            folder_input_dir.mkdir()
            (folder_input_dir / "sample.sql").write_text(sql, encoding="utf-8")

            folder_result = db_automation.process_folder(
                str(folder_input_dir),
                str(folder_output_dir),
                starting_seq=1,
                developer_name="Tester",
                description="Folder sample automation",
            )
            self.assertEqual(folder_result["errors"], [])
            folder_file = folder_output_dir / folder_result["processed"][0]["output"]
            folder_content = folder_file.read_text(encoding="utf-8")
            self.assertIn(
                "-- Description            :    Folder sample automation",
                folder_content,
            )

    def test_multi_schema_replaces_indented_use_line(self):
        content = "    USE db_old;\nSELECT 1;\n"
        converted = multi_schema_combiner.replace_use_statement(content, "db_new")
        self.assertNotIn("db_old", converted)
        self.assertEqual(converted.upper().count("USE "), 1)
        self.assertTrue(converted.startswith("USE db_new;"))

    def test_configured_schemas_flow_through_workfile_and_multi_schema(self):
        expected_schemas = multi_schema_combiner.DEFAULT_SCHEMAS
        self.assertEqual(len(expected_schemas), 7)

        header = workfile_generator.generate_header("PROCEDURE", "sample_sp", "1234")
        use_line = next(line for line in header.splitlines() if line.startswith("USE "))
        self.assertEqual(use_line, f"USE {expected_schemas[0]};")

    def test_workfile_to_multi_schema_uses_configured_defaults(self):
        raw_sql = (
            "CREATE PROCEDURE `sample_sp` ()\n"
            "BEGIN\n"
            "    SELECT 1;\n"
            "END"
        )
        workfile, _, _ = workfile_generator.process_sql(raw_sql, "1234")
        combined = multi_schema_combiner.combine_text_for_schemas_to_clipboard(
            workfile,
            multi_schema_combiner.DEFAULT_SCHEMAS,
        )

        self.assertEqual(combined["errors"], [])
        self.assertEqual(combined["schemas"], multi_schema_combiner.DEFAULT_SCHEMAS)
        expected_schemas = multi_schema_combiner.DEFAULT_SCHEMAS
        self.assertEqual(combined["content"].count("-- SCHEMA:"), len(expected_schemas))
        self.assertEqual(
            combined["content"].count("CREATE PROCEDURE `sample_sp`"), len(expected_schemas)
        )
        for schema in multi_schema_combiner.DEFAULT_SCHEMAS:
            self.assertEqual(combined["content"].count(f"USE {schema};"), 1)

    def test_workfile_generator_parses_algorithm_and_sql_security_view(self):
        view_sql = (
            "CREATE ALGORITHM=UNDEFINED DEFINER=`db_user`@`%` SQL SECURITY DEFINER "
            "VIEW `sample_first_task_vw` AS SELECT 1;"
        )
        sql_type, object_name = workfile_generator.detect_sql_type(view_sql)
        self.assertEqual(sql_type, "VIEW")
        self.assertEqual(object_name, "sample_first_task_vw")

        combined = view_sql + "\n\n" + view_sql.replace("`sample_first_task_vw`", "`app_window_vw`")
        parts = workfile_generator.split_sql_objects(combined)
        self.assertEqual(len(parts), 2)

        result = workfile_generator.process_pasted_content_to_clipboard(combined, "1234")
        self.assertEqual(len(result["processed"]), 2)
        self.assertEqual(result["errors"], [])

    def test_workfile_versions_increment_in_bracketed_jira_folder(self):
        raw_sql = "CREATE DEFINER=`db_user`@`%` PROCEDURE `sample_sp`()\nBEGIN\nSELECT 1;\nEND"
        with tempfile.TemporaryDirectory() as base_dir:
            folder = workfile_generator.create_jira_folder("1234", base_dir, "Bracketed [name]")
            first = workfile_generator.process_pasted_content(raw_sql, "1234", base_dir)
            second = workfile_generator.process_pasted_content(raw_sql, "1234", base_dir)
            self.assertEqual(first["folder_path"], folder)
            self.assertNotEqual(first["processed"][0]["main_file"], second["processed"][0]["main_file"])
            self.assertTrue(second["processed"][0]["main_file"].split(" ")[0].endswith("_02"))
            self.assertEqual(len(list(Path(folder).glob("*.sql"))), 2)
            self.assertEqual(len(list((Path(folder) / "Backup").glob("*.sql"))), 2)

    def test_db_automation_invalid_input_folder_returns_structured_error(self):
        with tempfile.TemporaryDirectory() as output_dir:
            missing = str(Path(output_dir) / "missing-folder")
            result = db_automation.process_folder(missing, output_dir)
            self.assertEqual(result["processed"], [])
            self.assertTrue(result["errors"])
            self.assertIn("Input folder not found", result["errors"][0])

    def test_db_automation_folder_processing_is_deterministic(self):
        with tempfile.TemporaryDirectory() as tmp:
            input_dir = Path(tmp) / "input"
            output_dir = Path(tmp) / "output"
            input_dir.mkdir()

            (input_dir / "b.sql").write_text(
                "CREATE FUNCTION `b_fn` ()\nRETURNS INT\nBEGIN\n    RETURN 1;\nEND;\n",
                encoding="utf-8",
            )
            (input_dir / "a.sql").write_text(
                "CREATE FUNCTION `a_fn` ()\nRETURNS INT\nBEGIN\n    RETURN 2;\nEND;\n",
                encoding="utf-8",
            )

            result = db_automation.process_folder(str(input_dir), str(output_dir))
            self.assertEqual(result["errors"], [])
            sources = [item["source"] for item in result["processed"]]
            self.assertEqual(sources, ["a.sql", "b.sql"])
            sequences = [item["sequence"] for item in result["processed"]]
            self.assertEqual(sequences, [1, 2])


# Miniature fixtures faithful to MySQL Workbench export formats
SQL_EXPORT_A = (
    "/*\n"
    "-- Query: SELECT * FROM schema_b.sample_items\n"
    "-- Date: 2026-07-15 23:12\n"
    "*/\n"
    "INSERT INTO `` (`item_id`,`item_name`,`metadata`,`updated_by`,`updated_time`) "
    "VALUES ('AddEntry','AddEntry','[{\\\"id\\\": 1, \\\"enabled\\\": 1}]',NULL,'2025-07-01 09:05:01');\n"
    "INSERT INTO `` (`item_id`,`item_name`,`metadata`,`updated_by`,`updated_time`) "
    "VALUES ('UpdateEntry','It''s a path',NULL,3,'2023-04-26 05:19:14');\n"
    "INSERT INTO `` (`item_id`,`item_name`,`metadata`,`updated_by`,`updated_time`) "
    "VALUES ('onlyInSourceA','onlyInSourceA','',NULL,NULL);\n"
)

CSV_EXPORT_B = (
    "item_id,item_name,metadata,updated_by,updated_time\n"
    'AddEntry,AddEntry,"[{"id": 1, "enabled": 2}]",NULL,"2025-07-02 10:00:00"\n'
    "UpdateEntry,It's a path,NULL,3,\"2023-04-26 05:19:14\"\n"
    "onlyInSourceB,onlyInSourceB,NULL,,NULL\n"
)

# CSV encoding of exactly the same data as SQL_EXPORT_A
CSV_EXPORT_A = (
    "item_id,item_name,metadata,updated_by,updated_time\n"
    'AddEntry,AddEntry,"[{"id": 1, "enabled": 1}]",NULL,"2025-07-01 09:05:01"\n'
    "UpdateEntry,It's a path,NULL,3,\"2023-04-26 05:19:14\"\n"
    "onlyInSourceA,onlyInSourceA,,NULL,NULL\n"
)


class TableComparatorTests(unittest.TestCase):
    def test_sql_parser_extracts_columns_rows_and_names(self):
        parsed = table_comparator.parse_sql_export(SQL_EXPORT_A)
        self.assertEqual(parsed["table_name"], "sample_items")
        self.assertEqual(parsed["schema"], "schema_b")
        self.assertEqual(
            parsed["columns"],
            ["item_id", "item_name", "metadata", "updated_by", "updated_time"],
        )
        self.assertEqual(len(parsed["rows"]), 3)
        self.assertEqual(parsed["warnings"], [])
        self.assertIsNone(parsed["rows"][0]["updated_by"])
        self.assertEqual(parsed["rows"][2]["metadata"], "")
        self.assertIsNone(parsed["rows"][2]["updated_time"])

    def test_sql_values_tokenizer_handles_escapes(self):
        values = table_comparator.split_sql_values_tuple(
            "('a,b','It''s','[{\\\"k\\\": 1}]',NULL,42,'')"
        )
        self.assertEqual(values, ["a,b", "It's", '[{"k": 1}]', None, "42", ""])

    def test_csv_parser_handles_embedded_json_quotes(self):
        parsed = table_comparator.parse_workbench_csv(CSV_EXPORT_B)
        self.assertEqual(len(parsed["rows"]), 3)
        self.assertEqual(parsed["warnings"], [])
        row = parsed["rows"][0]
        self.assertEqual(row["metadata"], '[{"id": 1, "enabled": 2}]')
        self.assertIsNone(row["updated_by"])
        self.assertEqual(row["updated_time"], "2025-07-02 10:00:00")
        self.assertEqual(parsed["rows"][2]["updated_by"], "")

    def test_csv_parser_handles_json_with_string_values(self):
        content = (
            "item_id,item_name_json,metadata\n"
            'OW_VIEW,"{"de": "Ansicht, gut", "en": "View"}","[{"id": 1, "enabled": 1}]"\n'
        )
        parsed = table_comparator.parse_workbench_csv(content)
        self.assertEqual(parsed["warnings"], [])
        row = parsed["rows"][0]
        self.assertEqual(row["item_name_json"], '{"de": "Ansicht, gut", "en": "View"}')
        self.assertEqual(row["metadata"], '[{"id": 1, "enabled": 1}]')

    def test_csv_row_width_mismatch_warns_and_skips(self):
        content = (
            "item_id,item_name,metadata\n"
            "good,good,NULL\n"
            "bad,too,many,fields,here\n"
        )
        parsed = table_comparator.parse_workbench_csv(content)
        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(len(parsed["warnings"]), 1)
        self.assertIn("row skipped", parsed["warnings"][0])

    def test_format_detection(self):
        self.assertEqual(table_comparator.detect_format("export.sql", ""), "sql")
        self.assertEqual(table_comparator.detect_format("export.csv", ""), "csv")
        self.assertEqual(
            table_comparator.detect_format("export.txt", "INSERT INTO `` (`a`) VALUES (1);"),
            "sql",
        )
        self.assertEqual(table_comparator.detect_format("export.txt", "a,b\n1,2\n"), "csv")

    def test_infer_source_label(self):
        schema_a, schema_b, _schema_c = table_comparator.SCHEMA_CODES[:3]
        self.assertEqual(
            table_comparator.infer_source_label(f"dev{schema_a}.sql"),
            {"env": "DEV", "schema": schema_a.upper(), "label": f"DEV {schema_a.upper()}"},
        )
        self.assertEqual(
            table_comparator.infer_source_label(f"prod{schema_b}.csv")["label"],
            f"PROD {schema_b.upper()}",
        )
        self.assertEqual(
            table_comparator.infer_source_label(
                f"sample ( dev {schema_a.lower()} ).sql"
            )["label"],
            f"DEV {schema_a.upper()}",
        )
        parsed = table_comparator.parse_sql_export(SQL_EXPORT_A)
        parsed_schema = table_comparator.schema_code(parsed["schema"]).upper()
        self.assertEqual(
            table_comparator.infer_source_label("test export.sql", parsed)["label"],
            f"TEST {parsed_schema}",
        )
        self.assertEqual(
            table_comparator.infer_source_label("sample_export.csv")["label"],
            "sample_export",
        )

    def test_two_way_diff_statuses_and_counts(self):
        table_a = table_comparator.parse_sql_export(SQL_EXPORT_A)
        table_b = table_comparator.parse_workbench_csv(CSV_EXPORT_B)
        result = table_comparator.compare_tables(table_a, table_b, "item_id")
        self.assertEqual(
            result["counts"], {"identical": 1, "changed": 1, "only_a": 1, "only_b": 1}
        )
        changed = next(row for row in result["rows"] if row["status"] == "changed")
        self.assertEqual(changed["key"], "AddEntry")
        self.assertEqual(changed["changed_columns"], {"metadata", "updated_time"})

    def test_two_way_ignored_columns_make_rows_identical(self):
        table_a = table_comparator.parse_sql_export(SQL_EXPORT_A)
        table_b = table_comparator.parse_workbench_csv(CSV_EXPORT_B)
        result = table_comparator.compare_tables(
            table_a, table_b, "item_id",
            ignored_columns={"metadata", "updated_time"},
        )
        self.assertEqual(result["counts"]["changed"], 0)
        self.assertEqual(result["counts"]["identical"], 2)
        self.assertNotIn("metadata", result["columns"])
        self.assertIn("item_id", result["columns"])

    def test_two_way_alternate_key_and_duplicate_keys(self):
        table_a = table_comparator.parse_sql_export(SQL_EXPORT_A)
        table_b = table_comparator.parse_workbench_csv(CSV_EXPORT_B)
        result = table_comparator.compare_tables(table_a, table_b, "item_name")
        keys = {row["key"] for row in result["rows"]}
        self.assertIn("It's a path", keys)

        duplicated = {
            "table_name": None,
            "schema": None,
            "columns": ["id", "val"],
            "rows": [{"id": "1", "val": "x"}, {"id": "1", "val": "y"}],
            "warnings": [],
        }
        other = {
            "table_name": None,
            "schema": None,
            "columns": ["id", "val"],
            "rows": [{"id": "1", "val": "x"}],
            "warnings": [],
        }
        result = table_comparator.compare_tables(duplicated, other, "id")
        self.assertEqual(result["counts"]["identical"], 1)
        self.assertTrue(any("duplicate key" in warning for warning in result["warnings"]))

    def test_two_way_column_mismatch_and_missing_key(self):
        table_a = {
            "table_name": None, "schema": None,
            "columns": ["id", "a_only", "shared"],
            "rows": [{"id": "1", "a_only": "x", "shared": "s"}],
            "warnings": [],
        }
        table_b = {
            "table_name": None, "schema": None,
            "columns": ["id", "shared", "b_only"],
            "rows": [{"id": "1", "shared": "s", "b_only": "y"}],
            "warnings": [],
        }
        result = table_comparator.compare_tables(table_a, table_b, "id")
        self.assertEqual(result["columns"], ["id", "shared"])
        self.assertEqual(len(result["warnings"]), 2)
        with self.assertRaises(ValueError):
            table_comparator.compare_tables(table_a, table_b, "a_only")

    def test_null_vs_empty_semantics(self):
        table_a = {
            "table_name": None, "schema": None, "columns": ["id", "val"],
            "rows": [{"id": "1", "val": None}, {"id": "2", "val": None}], "warnings": [],
        }
        table_b = {
            "table_name": None, "schema": None, "columns": ["id", "val"],
            "rows": [{"id": "1", "val": None}, {"id": "2", "val": ""}], "warnings": [],
        }
        result = table_comparator.compare_tables(table_a, table_b, "id")
        self.assertEqual(result["counts"]["identical"], 1)
        self.assertEqual(result["counts"]["changed"], 1)
        self.assertEqual(table_comparator.format_cell(None), "NULL")

    def test_sql_and_csv_of_same_data_compare_identical(self):
        table_sql = table_comparator.parse_sql_export(SQL_EXPORT_A)
        table_csv = table_comparator.parse_workbench_csv(CSV_EXPORT_A)
        result = table_comparator.compare_tables(table_sql, table_csv, "item_id")
        self.assertEqual(
            result["counts"], {"identical": 3, "changed": 0, "only_a": 0, "only_b": 0}
        )

    def test_compare_multi_counts_and_statuses(self):
        table_a = table_comparator.parse_sql_export(SQL_EXPORT_A)
        table_b = table_comparator.parse_workbench_csv(CSV_EXPORT_B)
        table_c = table_comparator.parse_workbench_csv(CSV_EXPORT_A)
        schema_a, schema_b, schema_c = [code.upper() for code in table_comparator.SCHEMA_CODES[:3]]
        sources = [
            {"label": f"DEV {schema_a}", "table": table_a},
            {"label": f"TEST {schema_b}", "table": table_b},
            {"label": f"DEV {schema_c}", "table": table_c},
        ]
        result = table_comparator.compare_multi(sources, "item_id")
        self.assertEqual(
            result["sources"], [f"DEV {schema_a}", f"TEST {schema_b}", f"DEV {schema_c}"]
        )
        by_key = {row["key"]: row for row in result["rows"]}

        add_user = by_key["AddEntry"]
        self.assertEqual(add_user["status"], "data_diff")
        self.assertEqual(add_user["present_count"], 3)
        self.assertEqual(add_user["match_count"], 2)
        self.assertEqual(add_user["diff_columns"], {"metadata", "updated_time"})

        self.assertEqual(by_key["UpdateEntry"]["status"], "ok")
        self.assertEqual(by_key["onlyInSourceA"]["status"], "missing")
        self.assertEqual(by_key["onlyInSourceA"]["present_count"], 2)
        self.assertEqual(by_key["onlyInSourceB"]["present_count"], 1)

        ignored = table_comparator.compare_multi(
            sources, "item_id", ignored_columns={"metadata", "updated_time"}
        )
        self.assertEqual(
            {row["key"]: row["status"] for row in ignored["rows"]}["AddEntry"], "ok"
        )

    def test_compare_multi_rejects_duplicate_labels(self):
        table = table_comparator.parse_sql_export(SQL_EXPORT_A)
        label = f"DEV {table_comparator.SCHEMA_CODES[0].upper()}"
        sources = [
            {"label": label, "table": table},
            {"label": label, "table": table},
        ]
        with self.assertRaises(ValueError):
            table_comparator.compare_multi(sources, "item_id")

    def test_export_multi_to_excel(self):
        table_a = table_comparator.parse_sql_export(SQL_EXPORT_A)
        table_b = table_comparator.parse_workbench_csv(CSV_EXPORT_B)
        schema_a, schema_b = [code.upper() for code in table_comparator.SCHEMA_CODES[:2]]
        sources = [
            {"label": f"DEV {schema_a}", "table": table_a,
             "info": {"env": "DEV", "schema": schema_a, "label": f"DEV {schema_a}"}},
            {"label": f"TEST {schema_b}", "table": table_b,
             "info": {"env": "TEST", "schema": schema_b, "label": f"TEST {schema_b}"}},
        ]
        result = table_comparator.compare_multi(sources, "item_id")

        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            output = str(Path(tmp) / "comparison.xlsx")
            table_comparator.export_multi_to_excel(result, sources, output)
            sheet = load_workbook(output).active
            self.assertEqual(sheet.title, "sample_items")
            header = [cell.value for cell in sheet[1]]
            self.assertEqual(header[:2], ["db_env", "db_schema"])
            self.assertEqual(header[-2:], ["present_count", "data_match_count"])
            self.assertEqual(sheet.max_row, 1 + 3 + 3)
            first_row = [cell.value for cell in sheet[2]]
            self.assertEqual(first_row[0], "DEV")
            self.assertEqual(first_row[1], schema_a)
            self.assertEqual(first_row[2], "AddEntry")
            self.assertEqual(first_row[-2], 2)
            self.assertEqual(first_row[-1], 1)
            self.assertIn("NULL", first_row)


if __name__ == "__main__":
    unittest.main()
