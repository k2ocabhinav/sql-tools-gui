"""Build smoke checks: run each enabled feature's real service on synthetic data.

The packaged app runs this through ``--smoke-test`` so that trimming the bundle
(unused Qt/Python modules) is proven not to break real work such as writing an
Excel workbook, not just importing the pages.
"""

from __future__ import annotations

import tempfile
from collections.abc import Callable
from pathlib import Path


class _Context:
    job_id = "smoke-test"

    def check_cancelled(self) -> None:
        return None

    def report(self, _message: str, _value: int = -1) -> None:
        return None


def _insert(root: Path) -> None:
    from openpyxl import load_workbook

    from sqltools.services import commit_file_run, run_insert

    prepared = run_insert(
        _Context(),
        {
            "paste": (
                "/*\n-- Query: SELECT * FROM dev.sample\n*/\n"
                "INSERT INTO `` (`id`) VALUES (1);\nINSERT INTO `` (`id`) VALUES (2);"
            ),
            "files": [],
            "input_folder": "",
            "output_folder": str(root / "insert"),
            "date_prefix": "20260101",
            "generate_combined": True,
            "excel": True,
        },
    )
    if commit_file_run(_Context(), prepared).status != "success":
        raise RuntimeError("INSERT output was not committed")
    workbooks = list((root / "insert").glob("*.xlsx"))
    if not workbooks:
        raise RuntimeError("INSERT Excel summary was not written")
    load_workbook(workbooks[0], read_only=True).close()


def _db_automation(root: Path) -> None:
    from sqltools.services import commit_file_run, run_db_automation

    prepared = run_db_automation(
        _Context(),
        {
            "mode": "paste",
            "paste": "CREATE PROCEDURE `sample_sp`() BEGIN SELECT 1; END",
            "input_folder": "",
            "output_folder": str(root / "automation"),
            "sequence": 1,
            "developer": "Tester",
            "description": "Smoke test",
        },
    )
    if commit_file_run(_Context(), prepared).status != "success":
        raise RuntimeError("DB Automation output was not committed")


def _workfile(root: Path) -> None:
    from sqltools.services import commit_file_run, run_workfile

    prepared = run_workfile(
        _Context(),
        {
            "copy": False,
            "paste": "CREATE PROCEDURE `sample_sp`() BEGIN SELECT 1; END",
            "jira": "1234",
            "base_folder": str(root / "workfiles"),
            "developer": "Tester",
            "temp_prefix": "temp_",
            "description": "Smoke test",
            "create_jira_folder": True,
        },
    )
    if commit_file_run(_Context(), prepared).status != "success":
        raise RuntimeError("Workfile output was not committed")


def _multi_schema(root: Path) -> None:
    from sqltools.services import commit_file_run, run_multi_schema

    prepared = run_multi_schema(
        _Context(),
        {
            "copy": False,
            "mode": "paste",
            "paste": "USE old_schema;\nSELECT 1;",
            "files": [],
            "input_folder": "",
            "schemas": ["schema_a", "schema_b"],
            "output_folder": str(root / "schemas"),
            "filename": "release.sql",
        },
    )
    if commit_file_run(_Context(), prepared).status != "success":
        raise RuntimeError("Multi-Schema output was not committed")


def _table_compare(root: Path) -> None:
    from openpyxl import load_workbook

    from logic import table_comparator as tc

    left = tc.parse_workbench_csv_compact("id,name\n1,a\n2,b\n")
    right = tc.parse_workbench_csv_compact("id,name\n1,a\n3,c\n")
    result = tc.compare_compact([("A", left), ("B", right)], "id")
    target = root / "comparison.xlsx"
    tc.export_compact_multi_to_excel(result, str(target))
    load_workbook(target, read_only=True).close()


CHECKS: dict[str, Callable[[Path], None]] = {
    "INSERT_CONSOLIDATOR": _insert,
    "DB_AUTOMATION": _db_automation,
    "WORKFILE_GENERATOR": _workfile,
    "MULTI_SCHEMA": _multi_schema,
    "TABLE_COMPARE": _table_compare,
}


def run_feature_workflows(features: list[str]) -> str | None:
    """Return a description of the first failing workflow, or None when all pass."""
    with tempfile.TemporaryDirectory(prefix="sql-tools-smoke-") as tmp:
        root = Path(tmp)
        for feature in features:
            check = CHECKS.get(feature)
            if check is None:
                continue
            try:
                check(root)
            except Exception as error:  # noqa: BLE001 - report every failure kind
                return f"{feature}: {type(error).__name__}: {error}"
    return None
