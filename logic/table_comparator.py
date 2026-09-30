"""
Table Comparator Logic Module
=============================
Parses MySQL Workbench table exports (SQL INSERT dumps or CSV) and compares
the same table across two or more schemas/environments.

- 2-way compare: row-by-row diff keyed on a chosen column.
- N-way compare: per-key presence/match counts across all loaded sources
  (mirrors the manual Excel comparison workflow) plus an Excel export.

"""

import os
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import TextIO

from logic.insert_consolidator import extract_table_name
from logic.multi_schema_combiner import DEFAULT_SCHEMAS
from logic.private_defaults import schema_code

# Source filename suffixes are derived from the configured schema list.
SCHEMA_CODES = [schema_code(schema) for schema in DEFAULT_SCHEMAS]

ENVIRONMENTS = ["dev", "test", "prod"]

# One Workbench export statement per row:
#   INSERT INTO `` (`col1`,`col2`,...) VALUES (...);
# The table-name backticks are empty in Workbench exports, but accept a name too.
_INSERT_PATTERN = re.compile(
    r'INSERT\s+INTO\s+`[^`]*`\s*\(([^)]+)\)\s*VALUES\s*\((.*?)\);\s*(?:\r?\n|$)',
    re.IGNORECASE | re.DOTALL,
)

_SCHEMA_COMMENT_PATTERN = re.compile(
    r'--\s*Query:\s*SELECT\s+.+?\s+FROM\s+(\w+)\.\w+',
    re.IGNORECASE,
)


@dataclass(frozen=True, slots=True)
class TableData:
    """Compact immutable table representation used by the Qt application."""

    table_name: str | None
    schema: str | None
    columns: tuple[str, ...]
    rows: tuple[tuple[str | None, ...], ...]
    warnings: tuple[str, ...] = ()
    export_format: str = "csv"
    file_path: str = ""
    column_index: dict[str, int] = field(default_factory=dict, compare=False, repr=False)

    def __post_init__(self) -> None:
        object.__setattr__(self, "column_index", {name: i for i, name in enumerate(self.columns)})


@dataclass(frozen=True, slots=True)
class CompactCompareRow:
    key: str | None
    status: str
    source_rows: tuple[int | None, ...]
    present_count: int
    match_count: int
    diff_columns: tuple[int, ...]


@dataclass(frozen=True, slots=True)
class CompactComparison:
    sources: tuple[tuple[str, TableData], ...]
    columns: tuple[str, ...]
    key_column: str
    rows: tuple[CompactCompareRow, ...]
    counts: dict[str, int]
    warnings: tuple[str, ...]
    two_way: bool


def _statement_finished(statement: str) -> bool:
    """Return true when a Workbench INSERT ends at an unquoted semicolon."""
    quote = None
    escaped = False
    line_comment = False
    block_comment = False
    i = 0
    while i < len(statement):
        char = statement[i]
        nxt = statement[i + 1] if i + 1 < len(statement) else ""
        if line_comment:
            if char == "\n":
                line_comment = False
            i += 1
            continue
        if block_comment:
            if char == "*" and nxt == "/":
                block_comment = False
                i += 2
            else:
                i += 1
            continue
        if quote:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == quote:
                if nxt == quote:
                    i += 1
                else:
                    quote = None
            i += 1
            continue
        if char == "#" or (char == "-" and nxt == "-" and
                            (i + 2 == len(statement) or statement[i + 2].isspace())):
            line_comment = True
            i += 2 if char == "-" else 1
            continue
        if char == "/" and nxt == "*":
            block_comment = True
            i += 2
            continue
        if char in ("'", '"', "`"):
            quote = char
        elif char == ";":
            return not statement[i + 1:].strip()
        i += 1
    return False


def _iter_workbench_inserts(handle: TextIO, check_cancelled: Callable[[], None] | None = None):
    """Stream individual Workbench INSERT statements without splitting lines into a list."""
    start = re.compile(r"^\s*INSERT\s+INTO\b", re.IGNORECASE)
    header: list[str] = []
    header_chars = 0
    statement: list[str] | None = None
    for line_number, line in enumerate(handle, 1):
        if check_cancelled and line_number % 2048 == 0:
            check_cancelled()
        if statement is None:
            if start.match(line):
                statement = [line]
                if _statement_finished(line):
                    yield "".join(statement), line_number, "".join(header)
                    statement = None
            else:
                if header_chars + len(line) <= 1024 * 1024:
                    header.append(line)
                    header_chars += len(line)
        else:
            statement.append(line)
            if _statement_finished("".join(statement)):
                yield "".join(statement), line_number, "".join(header)
                statement = None
    if statement:
        yield None, line_number, "".join(header)


def parse_sql_export_compact(content: str) -> TableData:
    """Parse an in-memory SQL export into column tuples, not per-row dictionaries."""
    from io import StringIO

    return _parse_sql_stream(StringIO(content), "")


def _parse_sql_stream(
    handle: TextIO,
    file_path: str,
    check_cancelled: Callable[[], None] | None = None,
) -> TableData:
    warnings: list[str] = []
    columns: tuple[str, ...] | None = None
    rows: list[tuple[str | None, ...]] = []
    table_name = None
    schema = None
    insert_pattern = re.compile(
        r'INSERT\s+INTO\s+`[^`]*`\s*\(([^)]+)\)\s*VALUES\s*\((.*?)\)\s*;\s*$',
        re.IGNORECASE | re.DOTALL,
    )
    for statement, line_number, header in _iter_workbench_inserts(handle, check_cancelled):
        if header and table_name is None:
            table_name = extract_table_name(header)
            schema_match = _SCHEMA_COMMENT_PATTERN.search(header)
            schema = schema_match.group(1) if schema_match else None
        if statement is None:
            warnings.append(f"Line {line_number}: unterminated INSERT statement; source is incomplete.")
            break
        match = insert_pattern.search(statement.strip())
        if not match:
            warnings.append(f"Line {line_number}: unsupported INSERT syntax; statement skipped.")
            continue
        statement_columns = tuple(
            column.strip().strip("`") for column in match.group(1).split(",")
        )
        if columns is None:
            columns = statement_columns
            if any(not column for column in columns) or len(set(columns)) != len(columns):
                raise ValueError("SQL export has empty or duplicate column names.")
        elif statement_columns != columns:
            warnings.append(f"Line {line_number}: column list differs from the first INSERT; row skipped.")
            continue
        values = split_sql_values_tuple("(" + match.group(2) + ")")
        if len(values) != len(columns):
            warnings.append(
                f"Line {line_number}: expected {len(columns)} values, got {len(values)}; row skipped."
            )
            continue
        rows.append(tuple(values))
    if columns is None:
        raise ValueError("No supported INSERT statements found in SQL export.")
    return TableData(table_name, schema, columns, tuple(rows), tuple(warnings), "sql", file_path)


def parse_workbench_csv_compact(content: str) -> TableData:
    from io import StringIO

    return _parse_workbench_csv_stream(StringIO(content), "")


def _parse_workbench_csv_stream(
    handle: TextIO,
    file_path: str,
    check_cancelled: Callable[[], None] | None = None,
) -> TableData:
    lines = iter(handle)
    header = None
    header_line = 0
    for line_number, physical_line in enumerate(lines, 1):
        if physical_line.strip():
            header_line = line_number
            header = physical_line.rstrip("\r\n")
            break
    if header is None:
        raise ValueError("CSV export is empty.")
    columns = tuple(column.strip() for column in header.split(","))
    if any(not column for column in columns) or len(set(columns)) != len(columns):
        raise ValueError("CSV export has empty or duplicate column names.")
    expected = len(columns)
    rows: list[tuple[str | None, ...]] = []
    warnings: list[str] = []
    line_number = header_line
    while True:
        try:
            physical_line = next(lines)
        except StopIteration:
            break
        line_number += 1
        if check_cancelled and line_number % 2048 == 0:
            check_cancelled()
        line = physical_line.rstrip("\r\n")
        if not line.strip():
            continue
        record_start = line_number
        values, open_quote = split_csv_line(line)
        # Join continuation lines only for a syntactically open quoted field.
        while open_quote:
            try:
                continuation = next(lines)
            except StopIteration:
                warnings.append(f"Line {record_start}: unterminated quoted field; source is incomplete.")
                line = ""
                break
            line_number += 1
            if check_cancelled and line_number % 2048 == 0:
                check_cancelled()
            line += "\n" + continuation.rstrip("\r\n")
            values, open_quote = split_csv_line(line)
        if not line:
            continue
        if len(values) != expected:
            warnings.append(
                f"Line {record_start}: expected {expected} fields, got {len(values)}; row skipped."
            )
            continue
        rows.append(tuple(values))
    return TableData(None, None, columns, tuple(rows), tuple(warnings), "csv", file_path)


def load_export_file_compact(
    file_path: str,
    check_cancelled: Callable[[], None] | None = None,
) -> TableData:
    """Strict UTF-8 loader that streams CSV/Workbench INSERT records."""
    from itertools import chain

    path = Path(file_path)
    try:
        if check_cancelled:
            check_cancelled()
        with path.open("r", encoding="utf-8-sig", errors="strict", newline=None) as handle:
            first = handle.readline()
            rest = chain((first,), handle)
            export_format = detect_format(str(path), first)
            if export_format == "sql":
                # SQL metadata appears before the first INSERT; chain avoids buffering the file.
                return _parse_sql_stream(_ChainedTextReader(rest), str(path), check_cancelled)
            return _parse_workbench_csv_stream(_ChainedTextReader(rest), str(path), check_cancelled)
    except UnicodeDecodeError as error:
        raise ValueError(f"File is not valid UTF-8 near byte {error.start}.") from error


class _ChainedTextReader:
    """Minimal iterator adapter because parser helpers deliberately consume streams."""

    def __init__(self, iterator):
        self.iterator = iterator

    def __iter__(self):
        return self.iterator


def compact_to_legacy(table: TableData) -> dict:
    return {
        "table_name": table.table_name, "schema": table.schema,
        "columns": list(table.columns),
        "rows": [dict(zip(table.columns, row, strict=True)) for row in table.rows],
        "warnings": list(table.warnings), "format": table.export_format,
        "file_path": table.file_path,
    }


def compare_compact(
    sources: list[tuple[str, TableData]],
    key_column: str,
    included_columns: set[str] | None = None,
    check_cancelled: Callable[[], None] | None = None,
) -> CompactComparison:
    """Compare compact sources and retain only row offsets in the result."""
    if len(sources) < 2:
        raise ValueError("At least two sources are required for comparison.")
    labels = [label for label, _ in sources]
    if len(set(labels)) != len(labels):
        raise ValueError("Source labels must be unique.")
    if any(key_column not in table.column_index for _, table in sources):
        raise ValueError(f"Key column '{key_column}' is missing from a source.")
    compared_columns, column_warnings = _common_columns(
        [list(table.columns) for _, table in sources], set(), key_column
    )
    if included_columns is not None:
        compared_columns = [
            col for col in compared_columns if col == key_column or col in included_columns
        ]
    column_indexes = [[table.column_index[col] for col in compared_columns] for _, table in sources]
    key_indexes = [table.column_index[key_column] for _, table in sources]
    warnings = list(column_warnings)
    indexed: list[dict[str | None, int]] = []
    for (label, table), key_index in zip(sources, key_indexes, strict=True):
        index: dict[str | None, int] = {}
        for row_number, row in enumerate(table.rows):
            if check_cancelled and row_number % 2048 == 0:
                check_cancelled()
            key = row[key_index]
            if key in index:
                warnings.append(f"{label}: duplicate key {format_cell(key)!r} (row {row_number + 1} ignored).")
            else:
                index[key] = row_number
        indexed.append(index)

    def unique_keys():
        seen: set[str | None] = set()
        for source_index in indexed:
            for key in source_index:
                if key not in seen:
                    seen.add(key)
                    yield key

    columns = tuple(compared_columns)
    key_position = columns.index(key_column)
    value_positions = [i for i in range(len(columns)) if i != key_position]
    counts = (
        {"identical": 0, "changed": 0, "only_a": 0, "only_b": 0}
        if len(sources) == 2 else
        {"ok": 0, "data_diff": 0, "missing": 0, "missing_and_diff": 0}
    )
    results: list[CompactCompareRow] = []
    for key_number, key in enumerate(unique_keys()):
        if check_cancelled and key_number % 2048 == 0:
            check_cancelled()
        row_offsets = tuple(index.get(key) for index in indexed)
        present = [source_i for source_i, row_i in enumerate(row_offsets) if row_i is not None]
        diff_positions: list[int] = []
        for column_position in value_positions:
            cell_values = {
                sources[source_i][1].rows[row_offsets[source_i]][column_indexes[source_i][column_position]]
                for source_i in present
            }
            if len(cell_values) > 1:
                diff_positions.append(column_position)
        present_count = len(present)
        signatures: dict[tuple[str | None, ...], int] = {}
        for source_i in present:
            row_i = row_offsets[source_i]
            signature = tuple(
                sources[source_i][1].rows[row_i][column_indexes[source_i][column_position]]
                for column_position in value_positions
            )
            signatures[signature] = signatures.get(signature, 0) + 1
        match_count = max(signatures.values(), default=0)

        if len(sources) == 2:
            if present_count == 1:
                status = "only_a" if row_offsets[0] is not None else "only_b"
            else:
                status = "changed" if diff_positions else "identical"
        else:
            missing, drift = present_count < len(sources), len(signatures) > 1
            status = ("missing_and_diff" if missing and drift else
                      "missing" if missing else "data_diff" if drift else "ok")
        counts[status] += 1
        results.append(
            CompactCompareRow(key, status, row_offsets, present_count, match_count, tuple(diff_positions))
        )
    for _table_label, table in sources:
        warnings.extend(table.warnings)
    return CompactComparison(
        tuple(sources), columns, key_column, tuple(results), counts,
        tuple(warnings), len(sources) == 2,
    )


def export_compact_multi_to_excel(
    result: CompactComparison,
    output_path: str,
    check_cancelled: Callable[[], None] | None = None,
) -> None:
    """Write the current N-way comparison snapshot using openpyxl's streaming mode."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    maximum_rows = 1_048_576
    maximum_columns = 16_384
    total_rows = sum(len(table.rows) for _, table in result.sources)
    if total_rows + 1 > maximum_rows:
        raise ValueError(f"Excel supports at most {maximum_rows - 1:,} exported data rows per sheet.")
    if len(result.columns) + 4 > maximum_columns:
        raise ValueError(f"Excel supports at most {maximum_columns - 4:,} compared columns.")

    source_stats: dict[str | None, CompactCompareRow] = {row.key: row for row in result.rows}
    table_name = next((table.table_name for _, table in result.sources if table.table_name), None)
    workbook = Workbook(write_only=True)
    sheet_title = re.sub(r"[\[\]:*?/\\]", "_", table_name or "comparison")[:31] or "comparison"
    sheet = workbook.create_sheet(sheet_title)

    def text_cell(value: str | None):
        cell = WriteOnlyCell(sheet, value="" if value is None else str(value))
        cell.data_type = "s"
        return cell

    sheet.append([
        text_cell(value)
        for value in ("db_env", "db_schema", *result.columns, "present_count", "data_match_count")
    ])
    key_indexes = [table.column_index[result.key_column] for _, table in result.sources]
    for source_index, (_label, table) in enumerate(result.sources):
        info = infer_source_label(table.file_path, {"schema": table.schema})
        env, schema = info.get("env", ""), info.get("schema", "")
        for row_number, row in enumerate(table.rows):
            if check_cancelled and row_number % 2048 == 0:
                check_cancelled()
            stats = source_stats.get(row[key_indexes[source_index]])
            values = [text_cell(env), text_cell(schema)]
            for column in result.columns:
                value = row[table.column_index[column]]
                values.append(text_cell(format_cell(value)))
            values.extend([stats.present_count if stats else 0, stats.match_count if stats else 0])
            sheet.append(values)
    workbook.save(output_path)


def format_cell(value: str | None) -> str:
    """Display form of a cell value (SQL NULL is stored as None)."""
    return 'NULL' if value is None else value


def detect_format(file_path: str, content: str) -> str:
    """Return 'sql' or 'csv' based on file extension, falling back to content."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.sql':
        return 'sql'
    if ext == '.csv':
        return 'csv'
    if re.search(r'INSERT\s+INTO', content, re.IGNORECASE):
        return 'sql'
    return 'csv'


# ============================================================================
# SQL EXPORT PARSING
# ============================================================================

_SQL_ESCAPES = {
    '"': '"',
    "'": "'",
    '\\': '\\',
    'n': '\n',
    'r': '\r',
    't': '\t',
    '0': '\0',
    'b': '\b',
    'Z': '\x1a',
}


def split_sql_values_tuple(tuple_text: str) -> list:
    """
    Tokenize the inside of a VALUES tuple into Python values.

    Quoted strings honor MySQL backslash escapes and doubled '' quotes;
    bare NULL becomes None; other bare tokens (numbers) stay as strings.
    """
    text = tuple_text.strip()
    if text.startswith('(') and text.endswith(')'):
        text = text[1:-1]

    values = []
    token_chars = []
    was_quoted = False
    in_string = False
    i = 0
    length = len(text)

    def finalize():
        nonlocal token_chars, was_quoted
        raw = ''.join(token_chars)
        if was_quoted:
            values.append(raw)
        else:
            bare = raw.strip()
            values.append(None if bare.upper() == 'NULL' else bare)
        token_chars = []
        was_quoted = False

    while i < length:
        char = text[i]
        if in_string:
            if char == '\\' and i + 1 < length:
                escaped = text[i + 1]
                token_chars.append(_SQL_ESCAPES.get(escaped, escaped))
                i += 2
                continue
            if char == "'":
                if i + 1 < length and text[i + 1] == "'":
                    token_chars.append("'")
                    i += 2
                    continue
                in_string = False
                i += 1
                continue
            token_chars.append(char)
            i += 1
            continue
        if char == "'":
            if not "".join(token_chars).strip():
                token_chars.clear()
            in_string = True
            was_quoted = True
            i += 1
            continue
        if char == ',':
            finalize()
            i += 1
            continue
        if was_quoted and char.isspace():
            i += 1
            continue
        token_chars.append(char)
        i += 1

    finalize()
    return values


def parse_sql_export(content: str) -> dict:
    """Parse a MySQL Workbench SQL export into {table_name, schema, columns, rows, warnings}."""
    warnings = []
    table_name = extract_table_name(content)
    schema_match = _SCHEMA_COMMENT_PATTERN.search(content)
    schema = schema_match.group(1) if schema_match else None

    matches = _INSERT_PATTERN.findall(content)
    if not matches:
        raise ValueError("No INSERT statements found in SQL export.")

    columns = [col.strip().strip('`') for col in matches[0][0].split(',')]
    rows = []
    for index, (col_text, values_text) in enumerate(matches, 1):
        stmt_columns = [col.strip().strip('`') for col in col_text.split(',')]
        if stmt_columns != columns:
            warnings.append(f"Statement {index}: column list differs from the first INSERT - skipped.")
            continue
        values = split_sql_values_tuple('(' + values_text + ')')
        if len(values) != len(columns):
            warnings.append(
                f"Statement {index}: expected {len(columns)} values, got {len(values)} - skipped."
            )
            continue
        rows.append(dict(zip(columns, values, strict=True)))

    return {
        'table_name': table_name,
        'schema': schema,
        'columns': columns,
        'rows': rows,
        'warnings': warnings,
    }


# ============================================================================
# CSV EXPORT PARSING (Workbench CSV is not RFC-4180: embedded JSON keeps raw
# double quotes inside quoted fields, so a custom tokenizer is required)
# ============================================================================

def split_csv_line(line: str) -> tuple[list, bool]:
    """
    Tokenize one Workbench CSV line.

    A '"' closes a quoted field only when it is the last character or is
    followed by ','. For fields whose content starts with '{' or '[' (embedded
    JSON), the closing '"' must additionally be preceded by '}' or ']' so JSON
    string values like '"en": "View",' stay inside the field.
    NOTE: a JSON string value ending exactly in '}",' or ']",' would still
    mis-split; row-width validation in parse_workbench_csv surfaces that as a
    warning instead of producing a corrupt row.

    Returns (values, ended_inside_quotes).
    """
    values = []
    token_chars = []
    was_quoted = False
    in_quotes = False
    at_field_start = True
    i = 0
    length = len(line)

    def finalize():
        nonlocal token_chars, was_quoted
        raw = ''.join(token_chars)
        if was_quoted:
            values.append(raw)
        else:
            values.append(None if raw == 'NULL' else raw)
        token_chars = []
        was_quoted = False

    while i < length:
        char = line[i]
        if in_quotes:
            if char == '"':
                if i + 1 < length and line[i + 1] == '"':
                    token_chars.append('"')
                    i += 2
                    continue
                at_delimiter = i + 1 == length or line[i + 1] == ','
                json_field = bool(token_chars) and token_chars[0] in '{['
                json_closed = bool(token_chars) and token_chars[-1] in '}]'
                if at_delimiter and (not json_field or json_closed):
                    in_quotes = False
                    i += 1
                    continue
                token_chars.append('"')
                i += 1
                continue
            token_chars.append(char)
            i += 1
            continue
        if at_field_start and char == '"':
            in_quotes = True
            was_quoted = True
            at_field_start = False
            i += 1
            continue
        if char == ',':
            finalize()
            at_field_start = True
            i += 1
            continue
        token_chars.append(char)
        at_field_start = False
        i += 1

    finalize()
    return values, in_quotes


def parse_workbench_csv(content: str) -> dict:
    """Parse a MySQL Workbench CSV export into {table_name, schema, columns, rows, warnings}."""
    warnings = []
    lines = content.splitlines()

    header_index = next(
        (i for i, line in enumerate(lines) if line.strip()), None
    )
    if header_index is None:
        raise ValueError("CSV export is empty.")

    columns = [col.strip() for col in lines[header_index].split(',')]
    expected = len(columns)

    rows = []
    i = header_index + 1
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        values, open_quote = split_csv_line(line)
        joins = 0
        # A field with an embedded newline continues on the next physical line
        while (open_quote or len(values) < expected) and i + 1 < len(lines) and joins < 20:
            i += 1
            joins += 1
            line = line + '\n' + lines[i]
            values, open_quote = split_csv_line(line)
        if len(values) != expected:
            warnings.append(
                f"Line {i + 1}: expected {expected} fields, got {len(values)} - row skipped."
            )
            i += 1
            continue
        rows.append(dict(zip(columns, values, strict=True)))
        i += 1

    return {
        'table_name': None,
        'schema': None,
        'columns': columns,
        'rows': rows,
        'warnings': warnings,
    }


def load_export_file(file_path: str) -> dict:
    """Read and parse an export file (single entry point for the GUI)."""
    with open(file_path, encoding='utf-8', errors='replace') as handle:
        content = handle.read()
    export_format = detect_format(file_path, content)
    parsed = parse_sql_export(content) if export_format == 'sql' else parse_workbench_csv(content)
    parsed['format'] = export_format
    parsed['file_path'] = file_path
    return parsed


# ============================================================================
# SOURCE LABELING
# ============================================================================

def infer_source_label(file_path: str, parsed_table: dict | None = None) -> dict:
    """
    Infer {'env', 'schema', 'label'} from a filename like 'devA.sql' or
    'sample ( dev a ).sql'. Falls back to the schema in the SQL header
    comment, then to the filename stem.
    """
    stem = os.path.splitext(os.path.basename(file_path))[0]
    lowered = stem.lower()

    env = ''
    for candidate in ENVIRONMENTS:
        if candidate in lowered:
            env = candidate.upper()
            break

    # Strip env tokens so glued names like 'deva' still expose the code
    stripped = lowered
    for candidate in ENVIRONMENTS:
        stripped = stripped.replace(candidate, ' ')

    schema = ''
    for code in SCHEMA_CODES:
        if re.search(r'(?<![a-z])' + code + r'(?![a-z])', stripped):
            schema = code.upper()
            break
    if not schema and parsed_table and parsed_table.get('schema'):
        schema = schema_code(parsed_table['schema']).upper()

    if env and schema:
        label = f"{env} {schema}"
    elif env or schema:
        label = (env + ' ' + schema).strip()
    else:
        label = stem
    return {'env': env, 'schema': schema, 'label': label}


# ============================================================================
# COMPARISON
# ============================================================================

def _common_columns(column_lists: list, ignored_columns: set, key_column: str) -> tuple[list, list]:
    """Intersection of all column lists in the first list's order, minus ignored."""
    warnings = []
    first = column_lists[0]
    shared = set(first)
    for columns in column_lists[1:]:
        shared &= set(columns)
    for index, columns in enumerate(column_lists):
        extra = [col for col in columns if col not in shared]
        if extra:
            warnings.append(
                f"Source {index + 1}: columns not shared by all sources ignored: {', '.join(extra)}"
            )
    if not shared:
        raise ValueError("The files have no columns in common.")
    for index, columns in enumerate(column_lists):
        if key_column not in columns:
            raise ValueError(f"Key column '{key_column}' is missing from source {index + 1}.")
    compared = [
        col for col in first
        if col in shared and (col == key_column or col not in ignored_columns)
    ]
    if key_column not in compared:
        compared.insert(0, key_column)
    return compared, warnings


def _index_rows(table: dict, key_column: str, side_name: str, warnings: list) -> dict:
    """Map key value -> row dict, first occurrence wins on duplicates."""
    indexed = {}
    for position, row in enumerate(table['rows'], 1):
        key = row.get(key_column)
        if key in indexed:
            warnings.append(
                f"{side_name}: duplicate key {format_cell(key)!r} (row {position} ignored)."
            )
            continue
        indexed[key] = row
    return indexed


def compare_tables(table_a: dict, table_b: dict,
                   key_column: str,
                   ignored_columns: set | None = None) -> dict:
    """
    Two-way diff. Returns rows with status identical/changed/only_a/only_b,
    per-row changed_columns, and summary counts.
    """
    ignored = set(ignored_columns or ())
    columns, warnings = _common_columns(
        [table_a['columns'], table_b['columns']], ignored, key_column
    )

    rows_a = _index_rows(table_a, key_column, "Side A", warnings)
    rows_b = _index_rows(table_b, key_column, "Side B", warnings)

    result_rows = []
    counts = {'identical': 0, 'changed': 0, 'only_a': 0, 'only_b': 0}

    for key, row_a in rows_a.items():
        row_b = rows_b.get(key)
        if row_b is None:
            status, changed = 'only_a', set()
        else:
            changed = {
                col for col in columns
                if col != key_column and row_a.get(col) != row_b.get(col)
            }
            status = 'changed' if changed else 'identical'
        counts[status] += 1
        result_rows.append({
            'key': format_cell(key),
            'status': status,
            'a': row_a,
            'b': row_b,
            'changed_columns': changed,
        })

    for key, row_b in rows_b.items():
        if key in rows_a:
            continue
        counts['only_b'] += 1
        result_rows.append({
            'key': format_cell(key),
            'status': 'only_b',
            'a': None,
            'b': row_b,
            'changed_columns': set(),
        })

    return {
        'key_column': key_column,
        'columns': columns,
        'rows': result_rows,
        'counts': counts,
        'warnings': warnings,
    }


def compare_multi(sources: list, key_column: str,
                  ignored_columns: set | None = None) -> dict:
    """
    N-way comparison across sources = [{'label': str, 'table': dict}, ...].

    Per key: present_count (sources containing it), match_count (largest
    group of identical rows), status, diff_columns, per_source rows.
    """
    if len(sources) < 2:
        raise ValueError("At least two sources are required for comparison.")

    ignored = set(ignored_columns or ())
    labels = [source['label'] for source in sources]
    if len(set(labels)) != len(labels):
        raise ValueError("Source labels must be unique - rename the duplicates first.")
    columns, warnings = _common_columns(
        [source['table']['columns'] for source in sources], ignored, key_column
    )

    indexed = {}
    for source in sources:
        indexed[source['label']] = _index_rows(
            source['table'], key_column, source['label'], warnings
        )

    key_order = []
    seen_keys = set()
    for label in labels:
        for key in indexed[label]:
            if key not in seen_keys:
                seen_keys.add(key)
                key_order.append(key)

    value_columns = [col for col in columns if col != key_column]
    total = len(sources)
    result_rows = []
    counts = {'ok': 0, 'data_diff': 0, 'missing': 0, 'missing_and_diff': 0}

    for key in key_order:
        per_source = {label: indexed[label].get(key) for label in labels}
        present_rows = [row for row in per_source.values() if row is not None]
        present_count = len(present_rows)

        groups = {}
        for row in present_rows:
            signature = tuple(row.get(col) for col in value_columns)
            groups[signature] = groups.get(signature, 0) + 1
        match_count = max(groups.values()) if groups else 0

        diff_columns = {
            col for col in value_columns
            if len({row.get(col) for row in present_rows}) > 1
        }

        missing = present_count < total
        drifted = len(groups) > 1
        if missing and drifted:
            status = 'missing_and_diff'
        elif missing:
            status = 'missing'
        elif drifted:
            status = 'data_diff'
        else:
            status = 'ok'
        counts[status] += 1

        result_rows.append({
            'key': format_cell(key),
            'status': status,
            'present_count': present_count,
            'match_count': match_count,
            'diff_columns': diff_columns,
            'per_source': per_source,
        })

    return {
        'key_column': key_column,
        'columns': columns,
        'sources': labels,
        'rows': result_rows,
        'counts': counts,
        'total_sources': total,
        'warnings': warnings,
    }


# ============================================================================
# EXCEL EXPORT (stacked layout mirroring the manual comparison workbook)
# ============================================================================

def export_multi_to_excel(result: dict, sources: list, output_path: str) -> None:
    """
    Write an N-way comparison to an .xlsx: one row per (source, key), lead
    columns db_env / db_schema, the compared columns, then present_count and
    data_match_count for that key.
    """
    from openpyxl import Workbook

    key_column = result['key_column']
    columns = result['columns']
    stats_by_key = {row['key']: row for row in result['rows']}

    table_name = None
    for source in sources:
        table_name = source['table'].get('table_name') or table_name

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (table_name or 'comparison')[:31]

    header = ['db_env', 'db_schema'] + columns + ['present_count', 'data_match_count']
    sheet.append(header)

    for source in sources:
        info = source.get('info') or {}
        env = info.get('env', '')
        schema = info.get('schema', '')
        for row in source['table']['rows']:
            key = format_cell(row.get(key_column))
            stats = stats_by_key.get(key)
            if stats is None:
                continue
            sheet.append(
                [env, schema]
                + [format_cell(row.get(col)) for col in columns]
                + [stats['present_count'], stats['match_count']]
            )

    workbook.save(output_path)
